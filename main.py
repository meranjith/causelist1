import os
import re
import pdfplumber

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph
)


INPUT_FOLDER = "input"
OUTPUT_FILE = "Consolidated_Cause_List.pdf"
WORD_OUTPUT_FILE = "Consolidated_Cause_List.docx"


ADVOCATE_NAMES = [
    "MAHESH CHOWDHARY",
    #"A MAHESH CHOWDHARY",
    "NAGARAJA NAIDU",
    "ABHIMANYU",
    "KRISHIKA VAISHNAV",
    "SHAHBAAZ HUSSAIN",
]


styles = getSampleStyleSheet()

styles["BodyText"].leading = 11
styles["BodyText"].spaceBefore = 0
styles["BodyText"].spaceAfter = 0


records = []


def clean(text):

    if text is None:
        return ""

    text = str(text)
    text = text.replace("\n", " ")
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def is_italic_char(ch):
    """
    Detect whether a PDF character is italic/oblique.
    Different PDFs/fonts use different font names, so check
    several common indicators.
    """

    fontname = str(ch.get("fontname", "")).lower()

    return (
        "italic" in fontname
        or "oblique" in fontname
        or "slanted" in fontname
    )


def extract_italic_from_cell(page, cell_bbox):
    """
    Extract only italic text from a PDF table cell.

    cell_bbox:
        (x0, top, x1, bottom)
    """

    if not cell_bbox:
        return ""

    x0, top, x1, bottom = cell_bbox

    chars = []

    for ch in page.chars:

        cx0 = ch.get("x0", 0)
        cx1 = ch.get("x1", 0)
        ctop = ch.get("top", 0)
        cbottom = ch.get("bottom", 0)

        # Character must overlap the cell
        horizontal_overlap = (
            cx1 > x0 and
            cx0 < x1
        )

        vertical_overlap = (
            cbottom > top and
            ctop < bottom
        )

        if not horizontal_overlap or not vertical_overlap:
            continue

        if is_italic_char(ch):
            chars.append(ch)

    if not chars:
        return ""

    # Sort characters in reading order
    chars.sort(
        key=lambda c: (
            round(c.get("top", 0), 1),
            c.get("x0", 0)
        )
    )

    # Reconstruct lines based on vertical position
    lines = []

    current_line = []
    current_top = None

    for ch in chars:

        char_top = ch.get("top", 0)

        if current_top is None:
            current_top = char_top

        # New line
        elif abs(char_top - current_top) > 3:

            if current_line:
                lines.append(current_line)

            current_line = []
            current_top = char_top

        current_line.append(ch)

    if current_line:
        lines.append(current_line)

    result_lines = []

    for line_chars in lines:

        line_chars.sort(
            key=lambda c: c.get("x0", 0)
        )

        text = ""

        previous_x1 = None

        for ch in line_chars:

            char = ch.get("text", "")

            if not char:
                continue

            x0_char = ch.get("x0", 0)

            # Insert a space where there is a visible gap
            # between words.
            if (
                previous_x1 is not None
                and x0_char - previous_x1 > 2
                and not text.endswith(" ")
            ):
                text += " "

            text += char

            previous_x1 = ch.get("x1", x0_char)

        text = clean(text)

        if text:
            result_lines.append(text)

    return "\n".join(result_lines)


def get_advocate_name(filename):

    return (
        filename.replace(".pdf", "")
        .replace(".PDF", "")
        .replace("_", " ")
        .upper()
        .strip()
    )


def extract_case_number(text):

    text = str(text)

    m = re.search(
        r"\b[A-Z\.]{2,10}\s+\d+/\d+\b",
        text
    )

    if m:
        return m.group(0)

    return text.split("\n")[0].strip()


def strip_advocates(value):

    value = clean(value)

    for adv in ADVOCATE_NAMES:

        value = re.sub(
            re.escape(adv),
            "",
            value,
            flags=re.I
        )

    return clean(value)


def extract_petitioner(text):

    if not text:
        return ""

    text = str(text)

    text = text.replace("PET:", "")

    lines = []

    for line in text.split("\n"):

        line = clean(line)

        if not line:
            continue

        upper_line = line.upper()

        if "FOR P" in upper_line:
            break

        if "FOR R" in upper_line:
            break

        if "ADVOCATE" in upper_line:
            break

        if "AGA" in upper_line:
            break

        if "HCGP" in upper_line:
            break

        if any(
            adv in upper_line
            for adv in ADVOCATE_NAMES
        ):
            break

        lines.append(line)

    result = clean(" ".join(lines))

    return strip_advocates(result)


def extract_respondent(text):

    if not text:
        return ""

    text = str(text)

    text = text.replace("RES:", "")

    lines = []

    stop_words = [
        "AGA",
        "HCGP",
        "FOR R",
        "FOR P",
        "ADVOCATE",
        "NOTICE",
        "V/O",
        "SD",
        "GOVT ADVOCATE",
    ]

    for line in text.split("\n"):

        line = clean(line)

        if not line:
            continue

        upper_line = line.upper()

        if any(
            adv in upper_line
            for adv in ADVOCATE_NAMES
        ):
            break

        stop = False

        for word in stop_words:

            if word in upper_line:
                stop = True
                break

        if stop:
            break

        lines.append(line)

    result = clean(" ".join(lines))

    return strip_advocates(result)


def build_case_name(row):

    pet = row["petitioner"]
    res = row["respondent"]

    if row["bold_side"] == "PET":
        pet = f"<b>{pet}</b>"

    if row["bold_side"] == "RES":
        res = f"<b>{res}</b>"

    return Paragraph(
        f"{pet}<br/>vs<br/>{res}",
        styles["BodyText"]
    )


def process_pdf(pdf_path, advocate):

    print("Processing:", pdf_path)

    current_judge = ""
    current_status = ""
    current_ch = ""
    current_list = ""

    with pdfplumber.open(pdf_path) as pdf:

        for page in pdf.pages:

            table_objects = page.find_tables()

            if not table_objects:
                continue

            for table_obj in table_objects:

                table = table_obj.extract()

                if not table:
                    continue

                for row_index, row in enumerate(table):

                    if not row:
                        continue

                    row_text = " ".join(
                        str(x)
                        for x in row
                        if x
                    )

                    upper_text = row_text.upper()

                    if "THE HON" in upper_text:
                        current_judge = clean(row_text)

                    if "COURT HALL NO" in upper_text:

                        ch_match = re.search(
                            r"COURT HALL NO\s*:\s*(\d+)",
                            row_text,
                            re.I
                        )

                        list_match = re.search(
                            r"CAUSE LIST NO\.?\s*(\d+)",
                            row_text,
                            re.I
                        )

                        if ch_match:
                            current_ch = ch_match.group(1)

                        if list_match:
                            current_list = list_match.group(1)

                    status_keywords = [
                        "PRELIMINARY HEARING",
                        "ADMISSION",
                        "ORDERS",
                        "FURTHER HEARING",
                        "HEARING -",
                    ]

                    for keyword in status_keywords:

                        if keyword in upper_text:
                            current_status = clean(row_text)

                    if len(row) < 6:
                        continue

                    sl_no = clean(row[0])

                    if not sl_no.isdigit():
                        continue

                    case_number = extract_case_number(
                        row[1]
                    )

                    pet_col = row[3] if len(row) > 3 else ""
                    res_col = row[5] if len(row) > 5 else ""

                    # --------------------------------------------------
                    # GET ACTUAL PETITIONER / RESPONDENT FROM ITALICS
                    # --------------------------------------------------

                    pet_bbox = None
                    res_bbox = None

                    if row_index < len(table_obj.rows):

                        cells = table_obj.rows[row_index].cells

                        if cells:

                            if len(cells) > 3:
                                pet_bbox = cells[3]

                            if len(cells) > 5:
                                res_bbox = cells[5]

                    petitioner = extract_italic_from_cell(
                        page,
                        pet_bbox
                    )

                    respondent = extract_italic_from_cell(
                        page,
                        res_bbox
                    )

                    # Fallback to old extraction
                    if not petitioner:
                        petitioner = extract_petitioner(
                            pet_col
                        )

                    if not respondent:
                        respondent = extract_respondent(
                            res_col
                        )

                    # Remove PET:/RES:
                    if petitioner:
                        petitioner = re.sub(
                            r"^\s*PET\s*:\s*",
                            "",
                            petitioner,
                            flags=re.I
                        )

                    if respondent:
                        respondent = re.sub(
                            r"^\s*RES\s*:\s*",
                            "",
                            respondent,
                            flags=re.I
                        )

                    petitioner = strip_advocates(
                        petitioner
                    )

                    respondent = strip_advocates(
                        respondent
                    )

                    # --------------------------------------------------
                    # DETERMINE WHICH SIDE THE ADVOCATE IS ON
                    # --------------------------------------------------

                    bold_side = ""

                    pet_text = str(pet_col).upper()
                    res_text = str(res_col).upper()
                    row_text_upper = row_text.upper()

                    adv_tokens = advocate.upper().split()

                    pet_match = all(
                        token in pet_text
                        for token in adv_tokens
                    )

                    res_match = all(
                        token in res_text
                        for token in adv_tokens
                    )

                    if pet_match:

                        bold_side = "PET"

                    elif res_match:

                        bold_side = "RES"

                    else:

                        advocate_found = all(
                            token in row_text_upper
                            for token in adv_tokens
                        )

                        if advocate_found:

                            pet_pos = row_text_upper.find("PET:")
                            res_pos = row_text_upper.find("RES:")

                            mahesh_pos = max(
                                row_text_upper.find(token)
                                for token in adv_tokens
                            )

                            if (
                                pet_pos != -1
                                and mahesh_pos > pet_pos
                                and (
                                    res_pos == -1
                                    or mahesh_pos < res_pos
                                )
                            ):
                                bold_side = "PET"

                            elif (
                                res_pos != -1
                                and mahesh_pos > res_pos
                            ):
                                bold_side = "RES"

                    # Debug specific case
                    if case_number == "WP 20069/2021":

                        print("CASE =", case_number)
                        print("ADVOCATE =", advocate.upper())
                        print("PET COL =", pet_col)
                        print("RES COL =", res_col)

                        print(
                            "ITALIC PET =",
                            petitioner
                        )

                        print(
                            "ITALIC RES =",
                            respondent
                        )

                        print(
                            "BOLD SIDE =",
                            bold_side
                        )

                    records.append({
                        "sl_no": sl_no,
                        "case_number": case_number,
                        "petitioner": petitioner,
                        "respondent": respondent,
                        "bold_side": bold_side,
                        "ch": current_ch,
                        "list": current_list,
                        "status": current_status,
                        "judge": current_judge,
                    })


def generate_pdf():

    data = [[
        "SL NO",
        "CASE NUMBER",
        "CASE NAME",
        "CH",
        "LIST",
        "SL NO",
        "STATUS",
        "JUDGES"
    ]]

    for idx, row in enumerate(records, start=1):

        data.append([
            str(idx),

            Paragraph(
                f"<b>{row['case_number']}</b>",
                styles["BodyText"]
            ),

            build_case_name(row),

            row["ch"],
            row["list"],
            row["sl_no"],

            Paragraph(
                row["status"],
                styles["BodyText"]
            ),

            Paragraph(
                row["judge"],
                styles["BodyText"]
            )
        ])

    doc = SimpleDocTemplate(
        OUTPUT_FILE,
        pagesize=landscape(A4),
        leftMargin=10,
        rightMargin=10,
        topMargin=10,
        bottomMargin=10
    )

    table = Table(
        data,
        repeatRows=1,
        colWidths=[
            35,
            90,
            220,
            35,
            35,
            35,
            150,
            170
        ]
    )

    table.setStyle(
        TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ])
    )

    doc.build([table])


# ============================================================
# WORD DOCUMENT FORMATTING HELPERS
# ============================================================

def set_cell_shading(cell, fill):

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = tc_pr.find(qn("w:shd"))

    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(qn("w:fill"), fill)


def set_cell_margins(
    cell,
    top=70,
    start=90,
    bottom=70,
    end=90
):

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_mar = tc_pr.first_child_found_in("w:tcMar")

    if tc_mar is None:

        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for m, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):

        node = tc_mar.find(qn(f"w:{m}"))

        if node is None:

            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)

        node.set(
            qn("w:w"),
            str(value)
        )

        node.set(
            qn("w:type"),
            "dxa"
        )


def set_cell_border(cell, **kwargs):

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_borders = tc_pr.first_child_found_in(
        "w:tcBorders"
    )

    if tc_borders is None:

        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in (
        "top",
        "start",
        "bottom",
        "end",
        "insideH",
        "insideV"
    ):

        if edge not in kwargs:
            continue

        edge_data = kwargs.get(edge)

        tag = "w:{}".format(edge)

        element = tc_borders.find(
            qn(tag)
        )

        if element is None:

            element = OxmlElement(tag)
            tc_borders.append(element)

        for key in [
            "val",
            "sz",
            "space",
            "color"
        ]:

            if key in edge_data:

                element.set(
                    qn("w:{}".format(key)),
                    str(edge_data[key])
                )


def set_table_fixed_layout(table):

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr

    tbl_layout = tbl_pr.find(
        qn("w:tblLayout")
    )

    if tbl_layout is None:

        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)

    tbl_layout.set(
        qn("w:type"),
        "fixed"
    )


def set_repeat_table_header(row):

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement(
        "w:tblHeader"
    )

    tbl_header.set(
        qn("w:val"),
        "true"
    )

    tr_pr.append(tbl_header)


def prevent_row_split(row):

    from docx.oxml import OxmlElement

    tr_pr = row._tr.get_or_add_trPr()

    cant_split = OxmlElement(
        "w:cantSplit"
    )

    tr_pr.append(cant_split)


def set_cell_width(
    cell,
    width_inches
):

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    cell.width = Inches(
        width_inches
    )

    tc_pr = cell._tc.get_or_add_tcPr()

    tc_w = tc_pr.find(
        qn("w:tcW")
    )

    if tc_w is None:

        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)

    tc_w.set(
        qn("w:w"),
        str(
            int(
                width_inches * 1440
            )
        )
    )

    tc_w.set(
        qn("w:type"),
        "dxa"
    )


def format_table_paragraph(
    paragraph,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    before=0,
    after=0,
    line_spacing=1.0
):

    paragraph.alignment = alignment

    paragraph.paragraph_format.space_before = Pt(
        before
    )

    paragraph.paragraph_format.space_after = Pt(
        after
    )

    paragraph.paragraph_format.line_spacing = (
        line_spacing
    )

    paragraph.paragraph_format.first_line_indent = (
        Inches(0)
    )

    paragraph.paragraph_format.left_indent = (
        Inches(0)
    )

    paragraph.paragraph_format.right_indent = (
        Inches(0)
    )


def style_cell_text(
    cell,
    font_size=8.5,
    bold=False,
    alignment=WD_ALIGN_PARAGRAPH.LEFT
):

    for paragraph in cell.paragraphs:

        format_table_paragraph(
            paragraph,
            alignment=alignment
        )

        for run in paragraph.runs:

            run.font.name = "Aptos"
            run.font.size = Pt(
                font_size
            )

            run.bold = bold


# ============================================================
# PROFESSIONAL WORD DOCUMENT
# ============================================================

def generate_word():

    doc = Document()

    # ---------------------------------------------------------
    # PAGE SETUP
    # ---------------------------------------------------------

    section = doc.sections[0]

    section.orientation = (
        WD_ORIENT.LANDSCAPE
    )

    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)

    section.left_margin = Inches(0.30)
    section.right_margin = Inches(0.30)

    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)

    # ---------------------------------------------------------
    # TABLE HEADERS
    # ---------------------------------------------------------

    headers = [
        "SL NO",
        "CASE NUMBER",
        "CASE NAME",
        "CH",
        "LIST",
        "SL NO",
        "STATUS",
        "JUDGES"
    ]

    table = doc.add_table(
        rows=1,
        cols=len(headers)
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.autofit = False

    table.style = "Table Grid"

    set_table_fixed_layout(
        table
    )

    # ---------------------------------------------------------
    # COLUMN WIDTHS
    # ---------------------------------------------------------

    widths = [

        0.45,   # SL NO

        1.20,   # CASE NUMBER

        3.10,   # CASE NAME

        0.45,   # CH

        0.50,   # LIST

        0.50,   # SL NO

        2.05,   # STATUS

        2.84,   # JUDGES
    ]

    # ---------------------------------------------------------
    # HEADER FORMATTING
    # ---------------------------------------------------------

    header_cells = (
        table.rows[0].cells
    )

    for i, header in enumerate(headers):

        cell = header_cells[i]

        cell.text = header

        set_cell_width(
            cell,
            widths[i]
        )

        set_cell_margins(
            cell,
            top=90,
            start=90,
            bottom=90,
            end=90
        )

        cell.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        # Professional blue header
        set_cell_shading(
            cell,
            "1F4E78"
        )

        set_cell_border(
            cell,

            top={
                "val": "single",
                "sz": "8",
                "color": "173A5E"
            },

            bottom={
                "val": "single",
                "sz": "8",
                "color": "173A5E"
            },

            start={
                "val": "single",
                "sz": "6",
                "color": "FFFFFF"
            },

            end={
                "val": "single",
                "sz": "6",
                "color": "FFFFFF"
            }
        )

        for paragraph in cell.paragraphs:

            format_table_paragraph(
                paragraph,
                alignment=(
                    WD_ALIGN_PARAGRAPH.CENTER
                )
            )

            for run in paragraph.runs:

                run.font.name = "Aptos"
                run.font.size = Pt(8)
                run.bold = True

    # Repeat header on every page
    set_repeat_table_header(
        table.rows[0]
    )

    # ---------------------------------------------------------
    # DATA ROWS
    # ---------------------------------------------------------

    for idx, row in enumerate(
        records,
        start=1
    ):

        cells = table.add_row().cells

        # Prevent case rows from splitting
        # awkwardly across pages.
        prevent_row_split(
            table.rows[-1]
        )

        values = [

            str(idx),

            row["case_number"],

            None,

            row["ch"],

            row["list"],

            row["sl_no"],

            row["status"],

            row["judge"]
        ]

        # -----------------------------------------------------
        # NORMAL CELLS
        # -----------------------------------------------------

        cells[0].text = values[0]
        cells[1].text = values[1]

        cells[3].text = values[3]
        cells[4].text = values[4]
        cells[5].text = values[5]

        cells[6].text = values[6]
        cells[7].text = values[7]

        # -----------------------------------------------------
        # CASE NAME
        # -----------------------------------------------------

        case_cell = cells[2]

        paragraph = (
            case_cell.paragraphs[0]
        )

        # Clear default paragraph
        paragraph.clear()

        pet = row["petitioner"]
        res = row["respondent"]

        # Petitioner
        run = paragraph.add_run(
            pet
        )

        run.font.name = "Aptos"
        run.font.size = Pt(8.5)

        if row["bold_side"] == "PET":

            run.bold = True

        # VS
        paragraph.add_run("\n")

        vs_run = paragraph.add_run(
            "vs"
        )

        vs_run.font.name = "Aptos"
        vs_run.font.size = Pt(7.5)
        vs_run.italic = True

        paragraph.add_run("\n")

        # Respondent
        run = paragraph.add_run(
            res
        )

        run.font.name = "Aptos"
        run.font.size = Pt(8.5)

        if row["bold_side"] == "RES":

            run.bold = True

        format_table_paragraph(
            paragraph,
            alignment=(
                WD_ALIGN_PARAGRAPH.LEFT
            ),
            before=0,
            after=0,
            line_spacing=1.0
        )

        # -----------------------------------------------------
        # FORMAT EVERY CELL
        # -----------------------------------------------------

        center_columns = {
            0,
            3,
            4,
            5
        }

        for i, cell in enumerate(cells):

            set_cell_width(
                cell,
                widths[i]
            )

            set_cell_margins(
                cell,
                top=80,
                start=90,
                bottom=80,
                end=90
            )

            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            # Alternating row shading
            if idx % 2 == 0:

                set_cell_shading(
                    cell,
                    "F5F8FB"
                )

            else:

                set_cell_shading(
                    cell,
                    "FFFFFF"
                )

            # Subtle professional borders
            set_cell_border(
                cell,

                top={
                    "val": "single",
                    "sz": "4",
                    "color": "D9E1E8"
                },

                bottom={
                    "val": "single",
                    "sz": "4",
                    "color": "D9E1E8"
                },

                start={
                    "val": "single",
                    "sz": "4",
                    "color": "D9E1E8"
                },

                end={
                    "val": "single",
                    "sz": "4",
                    "color": "D9E1E8"
                }
            )

            # Case name already formatted above
            if i == 2:
                continue

            if i in center_columns:

                alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                )

            else:

                alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                )

            style_cell_text(
                cell,
                font_size=8.5,
                bold=False,
                alignment=alignment
            )

        # -----------------------------------------------------
        # CASE NUMBER EMPHASIS
        # -----------------------------------------------------

        for run in (
            cells[1]
            .paragraphs[0]
            .runs
        ):

            run.bold = True
            run.font.name = "Aptos"
            run.font.size = Pt(8.5)

        # -----------------------------------------------------
        # JUDGE COLUMN
        # -----------------------------------------------------

        for paragraph in (
            cells[7].paragraphs
        ):

            format_table_paragraph(
                paragraph,
                alignment=(
                    WD_ALIGN_PARAGRAPH.LEFT
                ),
                before=0,
                after=0,
                line_spacing=1.0
            )

    # ---------------------------------------------------------
    # FINAL COLUMN WIDTH PASS
    # ---------------------------------------------------------

    for row in table.rows:

        for i, width in enumerate(widths):

            set_cell_width(
                row.cells[i],
                width
            )

    # ---------------------------------------------------------
    # DOCUMENT DEFAULTS
    # ---------------------------------------------------------

    if doc.paragraphs:

        for paragraph in doc.paragraphs:

            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)

    # ---------------------------------------------------------
    # SAVE
    # ---------------------------------------------------------

    doc.save(
        WORD_OUTPUT_FILE
    )

    print(
        "WORD CREATED:",
        WORD_OUTPUT_FILE
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print("STARTING")

    if not os.path.exists(
        INPUT_FOLDER
    ):

        print(
            "input folder not found"
        )

        return

    files = [

        f

        for f in os.listdir(
            INPUT_FOLDER
        )

        if f.lower().endswith(
            ".pdf"
        )
    ]

    print(
        "PDF COUNT:",
        len(files)
    )

    for file in files:

        advocate = get_advocate_name(
            file
        )

        path = os.path.join(
            INPUT_FOLDER,
            file
        )

        process_pdf(
            path,
            advocate
        )

    print(
        "RECORDS:",
        len(records)
    )

    if len(records) == 0:

        print(
            "NO RECORDS FOUND"
        )

        return

    # Sort all records by Court Hall
    records.sort(
        key=lambda row:
            int(row["ch"])
            if str(row["ch"]).isdigit()
            else 9999
    )

    # Generate PDF
    generate_pdf()

    # Generate Word document
    generate_word()

    print(
        "PDF CREATED:",
        OUTPUT_FILE
    )

    print(
        "WORD CREATED:",
        WORD_OUTPUT_FILE
    )


if __name__ == "__main__":
    main()
